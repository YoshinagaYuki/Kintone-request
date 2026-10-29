"""Read-only extraction of the original7Tezukuru tabs. Never writes the workbook."""
import sys,json,re,hashlib,unicodedata,calendar,datetime,uuid
import openpyxl
src,dest=sys.argv[1:3]
w=openpyxl.load_workbook(src,data_only=True)
tabs=['シールLAB','バウンドボールLAB','スクイーズLAB','実験ラボシーズン','てずくーる','スクイーズLAB(ベルパーク様)','てずくーる！！フェス']
norm=lambda v:unicodedata.normalize('NFKC',str(v or '')).strip()
entries=[];stocks=[];notes=[];coverage={};months={}
for tab in tabs:
 s=w[tab];merged={}
 for m in s.merged_cells.ranges:
  value=s.cell(m.min_row,m.min_col).value
  for rr in range(m.min_row,m.max_row+1):
   for cc in range(m.min_col,m.max_col+1):merged[(rr,cc)]=(value,s.cell(m.min_row,m.min_col).coordinate)
 def cell(r,c):return merged.get((r,c),(s.cell(r,c).value,s.cell(r,c).coordinate))
 month=None;monthnums={};last_staff=None;used=set();sourcecells=set()
 item=str(uuid.uuid5(uuid.NAMESPACE_URL,'tezukuru-inventory:'+tab))
 for r in range(1,s.max_row+1):
  a=norm(s.cell(r,1).value);m=re.match(r'^(\d{1,2})月',a)
  if m:month=int(m[1])
  if not month:continue
  if norm(s.cell(r,3).value)=='担当者':last_staff=r;continue
  isevent=norm(s.cell(r,3).value)=='イベント場所'
  orphan=not norm(s.cell(r,3).value) and last_staff is not None and r==last_staff+2 and any(norm(cell(r,c)[0]) for c in range(4,35))
  if not(isevent or orphan) or last_staff is None:continue
  lane=norm(cell(last_staff,2)[0]);
  if orphan:lane=''
  try:lane=str(int(float(lane)))
  except ValueError:notes.append({'tab':tab,'row':r,'reason':'番号未記載'});lane=''
  if lane:monthnums.setdefault(month,set()).add(int(lane))
  days=calendar.monthrange(2026,month)[1]
  vals=[norm(cell(r,d+3)[0]) for d in range(1,days+1)]
  refs=[cell(r,d+3)[1] for d in range(1,days+1)]
  # Keep each raw event anchor in provenance, even when an adjoining transport marker is folded in.
  for d,value in enumerate(vals):
   if value:sourcecells.add(refs[d])
  marker=lambda v: bool(re.search(r'[・_](?:納品|集荷|発送|返却)$',v)) or bool(re.fullmatch(r'(?:納品|集荷|配送中|返送中|返却|発送|移動|配送|予備|不可|戻り|使用不可|出荷|到着|回収|返却期間|\d+集荷)',v))
  groups=[];i=0
  while i<days:
   v=vals[i]
   if not v or marker(v):i+=1;continue
   end=i
   while end+1<days and vals[end+1]==v:end+=1
   groups.append((i,end,v));i=end+1
  owned=set()
  for gi,(a,b,value) in enumerate(groups):
   lo=groups[gi-1][1]+1 if gi else 0;hi=groups[gi+1][0]-1 if gi+1<len(groups) else days-1
   before=[d for d in range(max(lo,a-7),a) if vals[d] and marker(vals[d]) and re.search('納品|配送|発送|出荷|到着',vals[d])]
   after=[d for d in range(b+1,min(hi,b+7)+1) if vals[d] and marker(vals[d]) and re.search('集荷|配送|返却|回収|返送',vals[d])]
   prior_pickups=[d for d in range(lo,a) if re.search('集荷|回収|返却',vals[d])]
   if prior_pickups:before=[d for d in before if d>max(prior_pickups)]
   if gi+1<len(groups):
    next_start=groups[gi+1][0]
    after=[d for d in after if not(d>=next_start-7 and re.search('納品|配送|発送|出荷|到着',vals[d]))]
   delivery=[d for d in before if re.search('納品|到着',vals[d])];pickup=[d for d in after if re.search('集荷|回収|返却',vals[d])]
   start=min(before) if before else a;end=max(after) if after else b
   selected=range(start,end+1);owned.update(selected)
   staff=next((norm(cell(last_staff,d+4)[0]) for d in range(a,b+1) if norm(cell(last_staff,d+4)[0])),'担当者未設定')
   number=re.search(r'(?<!\d)(\d{3,6})(?!\d|個|名|台|枚|円)',re.sub(r'(?:通常版|数量|計|各)\d+','',value))
   date=lambda d:f'2026-{month:02d}-{d+1:02d}'
   cells=sorted(set(refs[d] for d in selected if vals[d]));key=hashlib.sha256(('tezukuru:'+tab+':'+','.join(cells)+':'+str(month)+':'+lane).encode()).hexdigest()
   raw={'cells':cells,'text':{refs[d]:vals[d] for d in selected if vals[d]},'usageStart':date(a),'usageEnd':date(b),'delivery':date(max(delivery) if delivery else max(before)) if before else '', 'pickup':date(min(pickup) if pickup else min(after)) if after else '', 'sourceRow':r,'annotation':bool(re.search('休業|祝日|年末年始|シルバーウィーク',value))}
   entries.append(dict(source_key=key,source_tab=tab,source_cell=','.join(cells),source_raw=json.dumps(raw,ensure_ascii=False),source_lane=lane,item_id=item,staff_name=staff,store_name=value,management_no=number[1] if number else '',start_date=date(start),end_date=date(end),request_id=None,state='skipped' if re.search('休業|祝日|年末年始|シルバーウィーク',value) else 'hold'))
  # Unattached shipping/blocked days remain explicit occupied source intervals.
  i=0
  while i<days:
   if i in owned or not vals[i]:i+=1;continue
   a=i;v=vals[i]
   while i+1<days and i+1 not in owned and vals[i+1]==v:i+=1
   b=i;owned.update(range(a,b+1));cells=sorted(set(refs[d] for d in range(a,b+1)));raw={'cells':cells,'text':v,'sourceRow':r,'transportOnly':True}
   key=hashlib.sha256(('tezukuru:'+tab+':'+','.join(cells)+':'+str(month)+':'+lane).encode()).hexdigest();num=re.search(r'(?<!\d)(\d{3,6})(?!\d|個|名|台|枚|円)',v)
   entries.append(dict(source_key=key,source_tab=tab,source_cell=','.join(cells),source_raw=json.dumps(raw,ensure_ascii=False),source_lane=lane,item_id=item,staff_name=norm(cell(last_staff,a+4)[0]) or '担当者未設定',store_name=v,management_no=num[1] if num else '',start_date=f'2026-{month:02d}-{a+1:02d}',end_date=f'2026-{month:02d}-{b+1:02d}',request_id=None,state='blocked'))
   i+=1
  used.update(refs[d] for d in owned if vals[d])
 # Future partial months must not silently reduce physical stock. Only a new maximum adds capacity.
 hist=max([max(v) for k,v in monthnums.items() if k<10]or[0]);current=max(monthnums.get(10,set())or monthnums.get(11,set())or {hist});total=max(hist,current)
 plans=[]
 if current<total:plans.append({'effective_date':'2026-10-01','total_count':current,'reserve_count':0})
 high=current
 for m in sorted(k for k in monthnums if k>10):
  count=max(monthnums[m]);
  if count>high:plans.append({'effective_date':f'2026-{m:02d}-01','total_count':count,'reserve_count':0});high=count
 stocks.append({'id':item,'name':tab,'category':'tezukuru','inventory_role':'rental_kit','aliases':[],'sort_order':2000+tabs.index(tab),'is_active':True,'total_count':total,'reserve_count':0,'plans':plans})
 coverage[tab]={'sourceCells':len(sourcecells),'coveredCells':len(used),'uncovered':sorted(sourcecells-used),'entries':sum(e['source_tab']==tab for e in entries),'numbersByMonth':{m:sorted(v) for m,v in monthnums.items()}}
# exact duplicates can come from vertical merges and are retained once, with stable source keys.
entries=list({e['source_key']:e for e in entries}.values())
source_data={name:{'cells':{c.coordinate:c.value for row in w[name].iter_rows()for c in row if c.value is not None},'merges':sorted(str(m)for m in w[name].merged_cells.ranges)}for name in tabs}
json.dump({'source':src,'stocks':stocks,'entries':entries,'coverage':coverage,'notes':notes,'source_data':source_data},open(dest,'w'),ensure_ascii=False,indent=2,default=str)
print(json.dumps({'entries':len(entries),'stock':[{'name':s['name'],'count':s['total_count'],'plans':s['plans']}for s in stocks],'uncovered':sum(len(v['uncovered'])for v in coverage.values()),'notes':notes},ensure_ascii=False))
